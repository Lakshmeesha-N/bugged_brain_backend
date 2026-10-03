import io
from docx import Document
from pypdf import PdfReader

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB


class UnsupportedFileType(Exception):
    """Raised when an uploaded file is not a supported type (.pdf or .docx)."""
    pass


class FileTooLarge(Exception):
    """Raised when an uploaded file exceeds the 5 MB limit."""
    pass


class NoTextFound(Exception):
    """Raised when the uploaded file contains fewer than 100 characters of readable text."""
    pass


def extract_text(filename: str, data: bytes) -> str:
    if len(data) > MAX_FILE_SIZE:
        raise FileTooLarge(f"File size exceeds 5 MB limit ({len(data)} bytes).")

    lower_filename = filename.lower()
    if not (lower_filename.endswith(".pdf") or lower_filename.endswith(".docx")):
        raise UnsupportedFileType(f"Unsupported file format for '{filename}'. Only .pdf and .docx are supported.")

    extracted_text = ""

    if lower_filename.endswith(".pdf"):
        try:
            reader = PdfReader(io.BytesIO(data))
            text_parts = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            extracted_text = "\n".join(text_parts)
        except Exception as e:
            raise UnsupportedFileType(f"Could not parse PDF file: {str(e)}") from e

    elif lower_filename.endswith(".docx"):
        try:
            doc = Document(io.BytesIO(data))
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text_parts.append(paragraph.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            text_parts.append(cell.text)
            extracted_text = "\n".join(text_parts)
        except Exception as e:
            raise UnsupportedFileType(f"Could not parse DOCX file: {str(e)}") from e

    cleaned_text = extracted_text.strip()
    if len(cleaned_text) < 100:
        raise NoTextFound("The file has no readable text, please upload a text-based PDF or DOCX.")

    return cleaned_text
